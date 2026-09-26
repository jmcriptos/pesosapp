# tests/test_reporte_kilos.py
"""Consulta «Kilos por lote»: totales por producto y lote sobre lo pesado.

La verdad de los kilos es la MISMA que la de `_kilos_y_cajas_pedido`: la
báscula manda; sin báscula, la línea de preparación; la línea original solo
cuenta si trae lote (datos históricos).
"""
import io
import os
from datetime import date, datetime, timedelta
from decimal import Decimal

import pytest

os.environ.setdefault('SECRET_KEY', 'test-secret')
os.environ.setdefault('FLASK_ENV', 'testing')
os.environ.setdefault('DATABASE_URL', 'sqlite:///:memory:')

from app import app as flask_app, db as _db

IDS = {}


@pytest.fixture
def app():
    flask_app.config.update(
        TESTING=True,
        WTF_CSRF_ENABLED=False,
        SQLALCHEMY_DATABASE_URI='sqlite:///:memory:',
    )
    with flask_app.app_context():
        _db.create_all()
        from app import (Rol, Territorio, Vendedor, Cliente, ClienteVendedor,
                         Producto, Pedido, DetallePedido, CajaPesada)

        rol_admin = Rol(nombre='super_admin', descripcion='Admin')
        rol_vend = Rol(nombre='vendedor', descripcion='Vendedor')
        terr = Territorio(nombre='t', descripcion='T')
        _db.session.add_all([rol_admin, rol_vend, terr])
        _db.session.flush()

        admin = Vendedor(username='admin', email='a@t.com', nombre_completo='Admin',
                         rol_id=rol_admin.id, territorio_id=terr.id, activo=True)
        admin.set_password('testpass')
        vend = Vendedor(username='vend', email='v@t.com', nombre_completo='Vendedor Uno',
                        rol_id=rol_vend.id, territorio_id=terr.id, activo=True)
        vend.set_password('testpass')
        _db.session.add_all([admin, vend])

        mangusa = Cliente(nombre='Mangusa', territorio_id=terr.id, qbo_id='C1')
        centrum = Cliente(nombre='Centrum', territorio_id=terr.id, qbo_id='C2')
        _db.session.add_all([mangusa, centrum])
        _db.session.flush()
        # El vendedor solo ve a Mangusa.
        _db.session.add(ClienteVendedor(cliente_id=mangusa.id, vendedor_id=vend.id, activo=True))

        chuleta = Producto(nombre='Chuleta de Cerdo', se_pesa=True, tax_rate=6.0, qbo_id='P1')
        chorizo = Producto(nombre='Chorizo Ahumado', se_pesa=True, tax_rate=6.0, qbo_id='P2')
        atun = Producto(nombre='Atún Van Camps', se_pesa=False, tax_rate=6.0, qbo_id='P3')
        _db.session.add_all([chuleta, chorizo, atun])
        _db.session.flush()

        hoy = datetime.utcnow()

        # PED 1 (Mangusa, facturado, hoy): chuleta con 3 cajas en dos lotes,
        # y una línea de preparación vieja del MISMO producto que NO debe
        # contar (la báscula manda).
        p1 = Pedido(cliente_id=mangusa.id, estado='facturado', fecha_pedido=hoy)
        _db.session.add(p1)
        _db.session.flush()
        d1 = DetallePedido(pedido_id=p1.id, producto_id=chuleta.id, cajas=3, cajas_pedidas=3,
                           peso=30, precio_unitario=10, subtotal=300, es_linea_pedido=True)
        _db.session.add(d1)
        _db.session.flush()
        _db.session.add_all([
            CajaPesada(detalle_pedido_id=d1.id, numero=1, peso=Decimal('10.500'), lote='L-0901',
                       fecha_elaboracion=date(2026, 9, 1), fecha_vencimiento=date(2026, 12, 1)),
            CajaPesada(detalle_pedido_id=d1.id, numero=2, peso=Decimal('9.250'), lote='L-0901',
                       fecha_elaboracion=date(2026, 9, 1), fecha_vencimiento=date(2026, 12, 1)),
            CajaPesada(detalle_pedido_id=d1.id, numero=3, peso=Decimal('11.000'), lote='L-0915',
                       fecha_elaboracion=date(2026, 9, 15), fecha_vencimiento=date(2026, 12, 15)),
        ])
        _db.session.add(DetallePedido(pedido_id=p1.id, producto_id=chuleta.id, cajas=1, cajas_pedidas=0,
                                      peso=99, lote='L-FANTASMA', precio_unitario=10, subtotal=990,
                                      es_linea_pedido=False))

        # PED 2 (Centrum, pendiente, hoy): chorizo pesado en el lote L-0915
        # (el mismo lote que la chuleta: los totales son por producto Y lote).
        p2 = Pedido(cliente_id=centrum.id, estado='pendiente', fecha_pedido=hoy)
        _db.session.add(p2)
        _db.session.flush()
        d2 = DetallePedido(pedido_id=p2.id, producto_id=chorizo.id, cajas=1, cajas_pedidas=1,
                           peso=5, precio_unitario=12, subtotal=60, es_linea_pedido=True)
        _db.session.add(d2)
        _db.session.flush()
        _db.session.add(CajaPesada(detalle_pedido_id=d2.id, numero=1, peso=Decimal('4.750'), lote='L-0915',
                                   fecha_elaboracion=date(2026, 9, 15), fecha_vencimiento=date(2026, 11, 15)))

        # PED 3 (Mangusa, entregado, hace 60 días): pedido anterior a la
        # báscula — la línea de preparación trae el peso y el lote. También
        # una línea de atún (por caja, sin peso) que suma cajas y no kilos.
        p3 = Pedido(cliente_id=mangusa.id, estado='entregado', fecha_pedido=hoy - timedelta(days=60))
        _db.session.add(p3)
        _db.session.flush()
        _db.session.add_all([
            DetallePedido(pedido_id=p3.id, producto_id=chuleta.id, cajas=2, cajas_pedidas=2, peso=20,
                          precio_unitario=10, subtotal=200, es_linea_pedido=True),
            DetallePedido(pedido_id=p3.id, producto_id=chuleta.id, cajas=2, cajas_pedidas=0, peso=21.5,
                          lote='L-0701', fecha_fabricacion='2026-07-01', fecha_expiracion='2026-10-01',
                          precio_unitario=10, subtotal=215, es_linea_pedido=False),
            DetallePedido(pedido_id=p3.id, producto_id=atun.id, cajas=4, cajas_pedidas=0, peso=0,
                          lote='AT-77', fecha_fabricacion='2026-01-01', fecha_expiracion='2028-01-01',
                          precio_unitario=50, subtotal=200, es_linea_pedido=False),
        ])

        # PED 4 (Centrum, pendiente, hoy): línea original SIN lote y sin
        # pesar — una promesa, no un kilo: no aparece.
        p4 = Pedido(cliente_id=centrum.id, estado='pendiente', fecha_pedido=hoy)
        _db.session.add(p4)
        _db.session.flush()
        _db.session.add(DetallePedido(pedido_id=p4.id, producto_id=chorizo.id, cajas=5, cajas_pedidas=5,
                                      peso=50, precio_unitario=12, subtotal=600, es_linea_pedido=True))
        _db.session.commit()
        IDS.update(mangusa=mangusa.id, centrum=centrum.id, chuleta=chuleta.id, chorizo=chorizo.id,
                   atun=atun.id, p1=p1.id, p2=p2.id, p3=p3.id, p4=p4.id)
        yield flask_app
        _db.drop_all()


def _login(app, username='admin'):
    client = app.test_client()
    client.post('/login', data={'username': username, 'password': 'testpass'}, follow_redirects=True)
    return client


def _filtros(**kw):
    from app import _leer_filtros_kilos
    from werkzeug.datastructures import MultiDict
    return _leer_filtros_kilos(MultiDict(kw))


def _lotes_de(grupos, producto):
    g = next(x for x in grupos if x['producto'] == producto)
    return {l['lote']: l for l in g['lotes']}


# ── Consulta ────────────────────────────────────────────────────────────────

def test_totaliza_por_producto_y_lote_con_la_bascula_mandando(app):
    from app import _consulta_kilos_por_lote
    with app.app_context():
        with app.test_request_context('/reportes/kilos?desde='):
            grupos, totales = _consulta_kilos_por_lote(_filtros(desde=''))
        chuleta = _lotes_de(grupos, 'Chuleta de Cerdo')
        assert set(chuleta) == {'L-0901', 'L-0915', 'L-0701'}
        assert chuleta['L-0901']['kg'] == Decimal('19.750')
        assert chuleta['L-0901']['cajas'] == 2
        assert chuleta['L-0915']['kg'] == Decimal('11.000')
        # La prep de 99 kg del PED 1 no cuenta: ese producto ya pasó por la báscula.
        assert 'L-FANTASMA' not in chuleta
        # Sin báscula, la línea de preparación es la verdad.
        assert chuleta['L-0701']['kg'] == Decimal('21.5')
        assert chuleta['L-0701']['fecha_elaboracion'] == date(2026, 7, 1)

        chorizo = _lotes_de(grupos, 'Chorizo Ahumado')
        assert set(chorizo) == {'L-0915'}          # el PED 4 (sin pesar, sin lote) no entra
        assert chorizo['L-0915']['kg'] == Decimal('4.750')

        atun = _lotes_de(grupos, 'Atún Van Camps')
        assert atun['AT-77']['cajas'] == 4 and atun['AT-77']['kg'] == 0

        assert totales['kg'] == Decimal('19.750') + Decimal('11.000') + Decimal('21.5') + Decimal('4.750')
        assert totales['pedidos'] == 3
        assert totales['lotes'] == 5


def test_el_mismo_lote_en_dos_productos_no_se_mezcla(app):
    from app import _consulta_kilos_por_lote
    with app.app_context(), app.test_request_context('/reportes/kilos'):
        grupos, _ = _consulta_kilos_por_lote(_filtros(lote='0915'))
        por_producto = {g['producto']: g['kg'] for g in grupos}
        assert por_producto == {'Chuleta de Cerdo': Decimal('11.000'), 'Chorizo Ahumado': Decimal('4.750')}
        # Y dentro de cada producto solo quedó el lote buscado.
        assert all(set(_lotes_de(grupos, p)) == {'L-0915'} for p in por_producto)


def test_filtra_por_producto_cliente_y_estado(app):
    from app import _consulta_kilos_por_lote
    with app.app_context(), app.test_request_context('/reportes/kilos'):
        grupos, totales = _consulta_kilos_por_lote(_filtros(producto_id=str(IDS['chuleta']), desde=''))
        assert [g['producto'] for g in grupos] == ['Chuleta de Cerdo']
        assert totales['kg'] == Decimal('19.750') + Decimal('11.000') + Decimal('21.5')

        grupos, totales = _consulta_kilos_por_lote(_filtros(cliente_id=str(IDS['centrum']), desde=''))
        assert [g['producto'] for g in grupos] == ['Chorizo Ahumado']

        grupos, totales = _consulta_kilos_por_lote(_filtros(estado='facturados', desde=''))
        assert totales['pedidos'] == 2 and 'Chorizo Ahumado' not in [g['producto'] for g in grupos]

        grupos, totales = _consulta_kilos_por_lote(_filtros(estado='activos', desde=''))
        assert [g['producto'] for g in grupos] == ['Chorizo Ahumado']


def test_sin_parametros_abre_en_los_ultimos_30_dias(app):
    from app import _consulta_kilos_por_lote
    with app.app_context(), app.test_request_context('/reportes/kilos'):
        filtros = _filtros()
        assert filtros['desde'] is not None and filtros['hasta'] is None
        grupos, _ = _consulta_kilos_por_lote(filtros)
        chuleta = _lotes_de(grupos, 'Chuleta de Cerdo')
        assert 'L-0701' not in chuleta                       # hace 60 días
        assert 'Atún Van Camps' not in [g['producto'] for g in grupos]

        # `desde` vacío mandado a propósito = todo el histórico.
        grupos, _ = _consulta_kilos_por_lote(_filtros(desde='', hasta=''))
        assert 'L-0701' in _lotes_de(grupos, 'Chuleta de Cerdo')


def test_el_vendedor_solo_ve_los_kilos_de_sus_clientes(app):
    from app import _consulta_kilos_por_lote
    client = _login(app, 'vend')
    with app.app_context():
        # La consulta corre dentro de una petición autenticada como vendedor.
        resp = client.get('/reportes/kilos?desde=&lote=0915')
        assert resp.status_code == 200
        html = resp.data.decode('utf-8')
        # El selector de producto lista el catálogo entero; lo que se acota
        # son los GRUPOS del reporte y el selector de clientes.
        assert '<h2>Chuleta de Cerdo</h2>' in html      # Mangusa: suyo
        assert '<h2>Chorizo Ahumado</h2>' not in html   # Centrum: no es suyo
        assert 'Centrum' not in html


# ── Pantalla y Excel ────────────────────────────────────────────────────────

def test_la_pantalla_muestra_totales_y_detalle(app):
    client = _login(app)
    resp = client.get('/reportes/kilos?desde=')
    assert resp.status_code == 200
    html = resp.data.decode('utf-8')
    assert 'Kilos por lote' in html
    assert 'L-0901' in html and 'L-0915' in html and 'L-0701' in html
    assert '19.75 kg' in html                       # subtotal del lote L-0901
    assert 'Total Chuleta de Cerdo' in html
    assert f'/pedidos/{IDS["p1"]}/detalles' in html  # el detalle enlaza al pedido
    assert 'L-FANTASMA' not in html
    # El enlace del Excel conserva la consulta.
    assert '/reportes/kilos/export?desde=' in html


def test_la_pantalla_vacia_no_revienta(app):
    client = _login(app)
    resp = client.get('/reportes/kilos?lote=NO-EXISTE')
    assert resp.status_code == 200
    assert 'No hay kilos registrados' in resp.data.decode('utf-8')


def test_el_excel_trae_resumen_y_detalle(app):
    from openpyxl import load_workbook
    client = _login(app)
    resp = client.get('/reportes/kilos/export?desde=&producto_id=%d' % IDS['chuleta'])
    assert resp.status_code == 200
    assert 'spreadsheetml' in resp.content_type
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == ['Resumen', 'Detalle']

    resumen = list(wb['Resumen'].iter_rows(values_only=True))
    filas = {r[1]: r for r in resumen if r[0] == 'Chuleta de Cerdo'}
    assert filas['L-0901'][5] == pytest.approx(19.75)
    assert filas['L-0901'][4] == 2
    assert filas['L-0701'][7] == 'Mangusa'
    total = next(r for r in resumen if r[0] == 'TOTAL GENERAL')
    assert total[5] == pytest.approx(19.75 + 11.0 + 21.5)

    detalle = list(wb['Detalle'].iter_rows(values_only=True))
    assert detalle[0][:2] == ('Fecha', 'Pedido')
    cajas_bascula = [r for r in detalle[1:] if r[11] == 'Báscula']
    assert len(cajas_bascula) == 3
    assert {r[1] for r in detalle[1:]} == {f'PED-{IDS["p1"]}', f'PED-{IDS["p3"]}'}


def test_el_reporte_esta_en_la_navegacion(app):
    client = _login(app)
    html = client.get('/pedidos').data.decode('utf-8')
    assert '/reportes/kilos' in html
    assert 'Kilos por lote' in html


def test_sin_sesion_redirige_al_login(app):
    resp = app.test_client().get('/reportes/kilos')
    assert resp.status_code in (302, 401)
