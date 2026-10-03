"""Producción propia: lotes, mermas y rendimiento SIN verificación de
disponibilidad de ingredientes."""
import os
from datetime import date
from decimal import Decimal

import pytest

os.environ.setdefault('SECRET_KEY', 'test-secret')
os.environ.setdefault('FLASK_ENV', 'testing')
os.environ.setdefault('DATABASE_URL', 'sqlite:///:memory:')

from app import app as flask_app, db as _db

IDS = {}


@pytest.fixture
def app():
    flask_app.config.update(TESTING=True, WTF_CSRF_ENABLED=False,
                            SQLALCHEMY_DATABASE_URI='sqlite:///:memory:')
    with flask_app.app_context():
        _db.create_all()
        from app import Rol, Territorio, Vendedor, Producto
        from maquila.models import Ingrediente, Receta, RecetaIngrediente
        ra = Rol(nombre='super_admin', descripcion='Admin')
        rv = Rol(nombre='vendedor', descripcion='Vendedor')
        terr = Territorio(nombre='t1', descripcion='T1')
        _db.session.add_all([ra, rv, terr])
        _db.session.flush()
        admin = Vendedor(username='admin', email='a@t.com', nombre_completo='Admin',
                         rol_id=ra.id, territorio_id=terr.id, activo=True)
        admin.set_password('pw')
        vend = Vendedor(username='vend', email='v@t.com', nombre_completo='Vend',
                        rol_id=rv.id, territorio_id=terr.id, activo=True)
        vend.set_password('pw')
        chorizo = Producto(nombre='Chorizo', se_pesa=True, tax_rate=10)
        jamon = Producto(nombre='Jamón', se_pesa=True, tax_rate=10)
        carne = Ingrediente(nombre='Carne de cerdo', unidad='kg')
        sal = Ingrediente(nombre='Sal', unidad='kg')
        tripa = Ingrediente(nombre='Tripa', unidad='ud')
        _db.session.add_all([admin, vend, chorizo, jamon, carne, sal, tripa])
        _db.session.flush()
        receta = Receta(producto_id=chorizo.id, cliente_id=None, nombre='Chorizo casa',
                        base_kg=Decimal('100'), activa=True)
        _db.session.add(receta)
        _db.session.flush()
        _db.session.add_all([
            RecetaIngrediente(receta_id=receta.id, ingrediente_id=carne.id, cantidad=Decimal('80')),
            RecetaIngrediente(receta_id=receta.id, ingrediente_id=sal.id, cantidad=Decimal('2')),
            RecetaIngrediente(receta_id=receta.id, ingrediente_id=tripa.id, cantidad=Decimal('50')),
        ])
        _db.session.commit()
        IDS.update(admin=admin.id, vend=vend.id, chorizo=chorizo.id, jamon=jamon.id,
                   carne=carne.id, sal=sal.id, tripa=tripa.id, receta=receta.id)
        yield flask_app
        _db.drop_all()


def _login(app, username):
    c = app.test_client()
    c.post('/login', data={'username': username, 'password': 'pw'}, follow_redirects=True)
    return c


def _lote(**kw):
    from produccion import servicios
    base = dict(producto_id=IDS['chorizo'], lote='L-1001',
                fecha_produccion=date(2026, 10, 1), vendedor_id=IDS['admin'],
                peso_producido=Decimal('85'),
                consumos={IDS['carne']: Decimal('100'), IDS['tripa']: Decimal('60')},
                mermas=[{'tipo': 'coccion', 'cantidad': Decimal('10'), 'motivo': ''}])
    base.update(kw)
    return servicios.crear_lote(**base)


# ---------------------------------------------------------------- andamiaje

def test_las_tablas_existen(app):
    with app.app_context():
        nombres = set(_db.inspect(_db.engine).get_table_names())
    assert {'lote_produccion', 'lote_consumo', 'lote_merma'} <= nombres


def test_las_tablas_se_crean_solas_si_faltan(app):
    """Heroku sin el script SQL: /produccion daba 500 por «no such table».
    Al arrancar, el módulo crea lo que falte, sin tocar lo que ya existe."""
    from produccion import asegurar_tablas
    from produccion.models import LoteProduccion, LoteConsumo, LoteMerma
    with app.app_context():
        for t in (LoteMerma.__table__, LoteConsumo.__table__, LoteProduccion.__table__):
            t.drop(_db.engine)
        assert 'lote_produccion' not in set(_db.inspect(_db.engine).get_table_names())
        asegurar_tablas(app)
        asegurar_tablas(app)   # segunda vez: no revienta por «ya existe»
        assert {'lote_produccion', 'lote_consumo', 'lote_merma'} <= set(
            _db.inspect(_db.engine).get_table_names())
    c = _login(app, 'admin')
    assert c.get('/produccion').status_code == 200
    assert c.get('/produccion/lotes').status_code == 200


def test_el_recurso_produccion_esta_en_los_permisos_configurables(app):
    from app import PERMISOS_RECURSOS, _permiso_default
    assert 'produccion' in PERMISOS_RECURSOS
    assert _permiso_default('vendedor', 'produccion', 'crear') is True
    assert _permiso_default('vendedor', 'produccion', 'editar') is False
    assert _permiso_default('supervisor', 'produccion', 'editar') is True


# ----------------------------------------------------------------- servicios

def test_registrar_sin_recepciones_ni_saldo_no_bloquea(app):
    """La regla que da nombre al módulo: no hay ledger ni recepciones, y el
    consumo se anota tal cual."""
    with app.app_context():
        from maquila.models import MovimientoIngrediente
        lote = _lote()
        assert lote.estado == 'abierta'
        assert lote.codigo == 'PR-2026-0001'
        assert {c.ingrediente_id: c.cantidad_real for c in lote.consumos} == {
            IDS['carne']: Decimal('100.000'), IDS['tripa']: Decimal('60.000')}
        assert MovimientoIngrediente.query.count() == 0


def test_el_codigo_es_correlativo_por_anio(app):
    with app.app_context():
        a = _lote(lote='A')
        b = _lote(lote='B')
        c = _lote(lote='C', fecha_produccion=date(2027, 1, 5))
        assert (a.codigo, b.codigo, c.codigo) == ('PR-2026-0001', 'PR-2026-0002', 'PR-2027-0001')


def test_balance_merma_identificada_y_sin_identificar(app):
    with app.app_context():
        from produccion import servicios
        b = servicios.balance(_lote())
        assert b['consumido'] == Decimal('100')          # la tripa (ud) no suma kilos
        assert b['otras_unidades'] == ['ud']
        assert b['producido'] == Decimal('85.000')
        assert b['merma_total'] == Decimal('15.000')
        assert b['merma_identificada'] == Decimal('10.000')
        assert b['merma_sin_identificar'] == Decimal('5.000')
        assert b['merma_pct'] == Decimal('15.0')
        assert b['rendimiento_pct'] == Decimal('85.0')
        assert b['merma_alta'] is True
        assert b['mermas_por_tipo'][0]['etiqueta'] == 'Cocción y ahumado'


def test_el_teorico_sale_de_la_receta_generica_y_queda_como_snapshot(app):
    with app.app_context():
        from produccion import servicios
        from maquila.models import RecetaIngrediente
        lote = _lote(peso_producido=Decimal('50'))
        assert lote.receta_id == IDS['receta']
        teoricos = {c.ingrediente_id: c.cantidad_teorica for c in lote.consumos}
        assert teoricos[IDS['carne']] == Decimal('40.000')
        assert teoricos[IDS['tripa']] == Decimal('25.000')
        # Cambiar la receta después no reescribe el lote.
        RecetaIngrediente.query.filter_by(receta_id=IDS['receta'],
                                          ingrediente_id=IDS['carne']).update({'cantidad': 90})
        _db.session.commit()
        b = servicios.balance(_db.session.get(type(lote), lote.id))
        carne = next(v for v in b['varianzas'] if v['ingrediente_id'] == IDS['carne'])
        assert carne['teorica'] == Decimal('40.000')
        assert carne['diferencia'] == Decimal('60.000')
        assert carne['pct'] == Decimal('150.0')


def test_producto_sin_receta_registra_solo_el_real(app):
    with app.app_context():
        from produccion import servicios
        lote = _lote(producto_id=IDS['jamon'], consumos={IDS['carne']: Decimal('20')},
                     peso_producido=Decimal('22'), mermas=[])
        assert lote.receta_id is None
        b = servicios.balance(lote)
        assert b['varianzas'][0]['teorica'] == Decimal('0')
        assert b['varianzas'][0]['pct'] is None
        assert b['merma_total'] == Decimal('-2.000')   # pesa más que lo que entró: se muestra
        assert b['rendimiento_pct'] == Decimal('110.0')


def test_lote_repetido_del_mismo_producto_se_rechaza_salvo_anulado(app):
    with app.app_context():
        from produccion import servicios
        primero = _lote()
        with pytest.raises(servicios.LoteInvalido):
            _lote()
        _lote(producto_id=IDS['jamon'], mermas=[])      # otro producto, mismo número: vale
        servicios.anular_lote(primero, IDS['admin'], 'tecleado dos veces')
        assert _lote().codigo == 'PR-2026-0003'


def test_validaciones_de_alta(app):
    with app.app_context():
        from produccion import servicios
        with pytest.raises(servicios.LoteInvalido):
            _lote(lote='   ')
        with pytest.raises(servicios.LoteInvalido):
            _lote(fecha_produccion=None)
        with pytest.raises(servicios.LoteInvalido):
            _lote(producto_id=9999)
        with pytest.raises(servicios.LoteInvalido):
            _lote(consumos={IDS['carne']: Decimal('-1')})
        with pytest.raises(servicios.LoteInvalido):
            _lote(mermas=[{'tipo': 'coccion', 'cantidad': Decimal('0'), 'motivo': 'x'}])
        with pytest.raises(servicios.LoteInvalido):
            _lote(mermas=[{'tipo': 'otro', 'cantidad': Decimal('1'), 'motivo': ''}])
        with pytest.raises(servicios.LoteInvalido):
            _lote(mermas=[{'tipo': 'inventado', 'cantidad': Decimal('1'), 'motivo': ''}])
        # Nada de lo rechazado quedó a medias.
        from produccion.models import LoteProduccion
        assert LoteProduccion.query.count() == 0


def test_consumo_cero_y_merma_vacia_se_omiten(app):
    with app.app_context():
        lote = _lote(consumos={IDS['carne']: Decimal('100'), IDS['sal']: Decimal('0')},
                     mermas=[{'tipo': 'coccion', 'cantidad': None, 'motivo': ''}])
        assert [c.ingrediente_id for c in lote.consumos] == [IDS['carne']]
        assert lote.mermas == []


def test_cerrar_exige_peso_y_consumo_en_kg(app):
    with app.app_context():
        from produccion import servicios
        sin_peso = _lote(lote='A', peso_producido=None)
        with pytest.raises(servicios.LoteInvalido):
            servicios.cerrar_lote(sin_peso, IDS['admin'])
        solo_tripa = _lote(lote='B', consumos={IDS['tripa']: Decimal('10')}, mermas=[])
        with pytest.raises(servicios.LoteInvalido):
            servicios.cerrar_lote(solo_tripa, IDS['admin'])
        ok = _lote(lote='C')
        servicios.cerrar_lote(ok, IDS['admin'])
        assert ok.estado == 'cerrada'
        assert ok.cerrado_por == IDS['admin']
        assert ok.cerrado_en is not None
        with pytest.raises(servicios.LoteInvalido):
            servicios.cerrar_lote(ok, IDS['admin'])


def test_editar_solo_abierto_y_reabrir_con_motivo(app):
    with app.app_context():
        from produccion import servicios
        lote = _lote()
        servicios.cerrar_lote(lote, IDS['admin'])
        cab = dict(producto_id=IDS['chorizo'], lote='L-1001', fecha_produccion=date(2026, 10, 1),
                   peso_producido=Decimal('90'))
        with pytest.raises(servicios.LoteNoEditable):
            servicios.editar_lote(lote, cabecera=cab, consumos={IDS['carne']: Decimal('100')})
        with pytest.raises(servicios.MotivoRequerido):
            servicios.reabrir_lote(lote, IDS['admin'], '  ')
        servicios.reabrir_lote(lote, IDS['admin'], 'se pesó mal la salida')
        assert lote.estado == 'abierta'
        assert lote.cerrado_en is None
        assert 'Reabierto: se pesó mal la salida' in lote.notas
        servicios.editar_lote(lote, cabecera=cab,
                              consumos={IDS['carne']: Decimal('100'), IDS['sal']: Decimal('2')},
                              mermas=[{'tipo': 'recorte', 'cantidad': Decimal('4'), 'motivo': ''}])
        assert lote.peso_producido == Decimal('90.000')
        assert {c.ingrediente_id for c in lote.consumos} == {IDS['carne'], IDS['sal']}
        # El teórico se recalcula con el peso nuevo.
        assert {c.ingrediente_id: c.cantidad_teorica for c in lote.consumos}[IDS['sal']] == Decimal('1.800')
        assert [m.tipo for m in lote.mermas] == ['recorte']
        b = servicios.balance(lote)
        assert b['consumido'] == Decimal('102')
        assert b['merma_sin_identificar'] == Decimal('8.000')


def test_editar_cambia_de_producto_y_resuelve_la_receta(app):
    with app.app_context():
        from produccion import servicios
        lote = _lote()
        cab = dict(producto_id=IDS['jamon'], lote='L-1001', fecha_produccion=date(2026, 10, 1),
                   peso_producido=Decimal('85'))
        servicios.editar_lote(lote, cabecera=cab, consumos={IDS['carne']: Decimal('100')})
        assert lote.producto_id == IDS['jamon']
        assert lote.receta_id is None
        assert lote.consumos[0].cantidad_teorica == Decimal('0')


def test_anular_exige_motivo_y_saca_de_los_reportes(app):
    with app.app_context():
        from produccion import servicios, reportes
        lote = _lote()
        servicios.cerrar_lote(lote, IDS['admin'])
        assert len(reportes.rendimiento()[0]) == 1
        with pytest.raises(servicios.MotivoRequerido):
            servicios.anular_lote(lote, IDS['admin'], '')
        servicios.anular_lote(lote, IDS['admin'], 'lote de prueba')
        assert lote.estado == 'anulada'
        assert lote.motivo_anulacion == 'lote de prueba'
        assert reportes.rendimiento()[0] == []
        with pytest.raises(servicios.LoteInvalido):
            servicios.anular_lote(lote, IDS['admin'], 'otra vez')


# ------------------------------------------------------------------ reportes

def _cerrados():
    from produccion import servicios
    a = _lote(lote='A', peso_producido=Decimal('85'),
              consumos={IDS['carne']: Decimal('100')},
              mermas=[{'tipo': 'coccion', 'cantidad': Decimal('10'), 'motivo': ''}])
    b = _lote(lote='B', fecha_produccion=date(2026, 10, 5), peso_producido=Decimal('190'),
              consumos={IDS['carne']: Decimal('200')},
              mermas=[{'tipo': 'coccion', 'cantidad': Decimal('6'), 'motivo': ''},
                      {'tipo': 'recorte', 'cantidad': Decimal('2'), 'motivo': ''}])
    j = _lote(lote='J', producto_id=IDS['jamon'], fecha_produccion=date(2026, 9, 20),
              peso_producido=Decimal('45'), consumos={IDS['carne']: Decimal('50')}, mermas=[])
    abierto = _lote(lote='X', peso_producido=Decimal('1'), consumos={IDS['carne']: Decimal('99')})
    for l in (a, b, j):
        servicios.cerrar_lote(l, IDS['admin'])
    return a, b, j, abierto


def test_rendimiento_pondera_por_kilos_y_filtra(app):
    with app.app_context():
        from produccion import reportes
        _cerrados()
        filas, resumen = reportes.rendimiento()
        assert [f['lote'] for f in filas] == ['B', 'A', 'J']       # cerrados, más reciente primero
        chorizo = next(r for r in resumen if r['producto'] == 'Chorizo')
        assert chorizo['lotes'] == 2
        assert chorizo['consumido'] == Decimal('300')
        assert chorizo['producido'] == Decimal('275.000')
        assert chorizo['rendimiento_pct'] == Decimal('91.7')        # 275/300, no (85+95)/2
        assert chorizo['merma_pct'] == Decimal('8.3')
        assert chorizo['merma_identificada'] == Decimal('18.000')
        assert chorizo['merma_sin_identificar'] == Decimal('7.000')
        assert (chorizo['rend_min'], chorizo['rend_max']) == (Decimal('85.0'), Decimal('95.0'))
        assert chorizo['lotes_merma_alta'] == 1

        filas, resumen = reportes.rendimiento(producto_id=IDS['jamon'])
        assert [f['lote'] for f in filas] == ['J']
        filas, _ = reportes.rendimiento(desde=date(2026, 10, 2))
        assert [f['lote'] for f in filas] == ['B']
        filas, _ = reportes.rendimiento(hasta=date(2026, 10, 2))
        assert [f['lote'] for f in filas] == ['A', 'J']


def test_mermas_por_causa_y_por_producto(app):
    with app.app_context():
        from produccion import reportes
        _cerrados()
        d = reportes.mermas()
        assert d['lotes'] == 3
        assert d['consumido'] == Decimal('350')
        assert d['merma_total'] == Decimal('30.000')
        assert d['identificada'] == Decimal('18.000')
        assert d['sin_identificar'] == Decimal('12.000')
        assert d['pct_identificada'] == Decimal('60.0')
        por_tipo = {t['tipo']: t for t in d['por_tipo']}
        assert set(por_tipo) == {'coccion', 'recorte', 'descarte', 'proceso', 'muestras', 'otro'}
        assert por_tipo['coccion']['cantidad'] == Decimal('16.000')
        assert por_tipo['coccion']['lotes'] == 2
        assert por_tipo['coccion']['pct_merma'] == Decimal('53.3')
        assert por_tipo['descarte']['cantidad'] == Decimal('0')
        assert d['por_tipo'][0]['tipo'] == 'coccion'
        jamon = next(p for p in d['por_producto'] if p['producto'] == 'Jamón')
        assert jamon['merma_total'] == Decimal('5.000')
        assert jamon['sin_identificar'] == Decimal('5.000')
        assert jamon['tipos'] == []


def test_resumen_del_periodo(app):
    with app.app_context():
        from produccion import reportes
        _cerrados()
        r = reportes.resumen_periodo(date(2026, 10, 1), date(2026, 10, 31))
        assert r['lotes'] == 2
        assert r['producido'] == Decimal('275.000')
        assert r['rendimiento_pct'] == Decimal('91.7')
        assert r['lotes_merma_alta'] == 1


# --------------------------------------------------------------------- rutas

def test_sin_sesion_redirige_al_login(app):
    r = app.test_client().get('/produccion', follow_redirects=False)
    assert r.status_code == 302


def test_admin_ve_el_resumen(app):
    c = _login(app, 'admin')
    assert c.get('/produccion').status_code == 200


def test_vendedor_lee_y_crea_pero_no_cierra(app):
    """Un solo cliente por test: la fixture deja un app context activo y los
    requests lo reusan, así que Flask-Login cachea en `g` el primer usuario
    que carga. Dos clientes en el mismo test verían al mismo usuario."""
    v = _login(app, 'vend')
    assert v.get('/produccion').status_code == 200              # leer por defecto
    assert v.get('/produccion/lotes/nuevo').status_code == 200  # crear por defecto
    with app.app_context():
        lote_id = _lote().id
    r = v.post(f'/produccion/lotes/{lote_id}/cerrar', follow_redirects=False)
    assert r.status_code == 302 and r.headers['Location'].endswith('/')  # editar: no
    r = v.get(f'/produccion/lotes/{lote_id}/editar', follow_redirects=False)
    assert r.status_code == 302 and r.headers['Location'].endswith('/')
    r = v.post(f'/produccion/lotes/{lote_id}/anular', data={'motivo': 'x'}, follow_redirects=False)
    assert r.status_code == 302 and r.headers['Location'].endswith('/')
    with app.app_context():
        from produccion.models import LoteProduccion
        assert _db.session.get(LoteProduccion, lote_id).estado == 'abierta'
    # Sin permiso de editar, «Guardar y cerrar» guarda y avisa; no cierra.
    r = v.post('/produccion/lotes/nuevo', data={
        'producto_id': str(IDS['chorizo']), 'lote': 'V-1',
        'fecha_produccion': '2026-10-02', 'peso_producido': '80',
        'consumo_ingrediente_id': [str(IDS['carne'])], 'consumo_real': ['100'],
        'accion': 'cerrar',
    }, follow_redirects=True)
    assert 'requiere permiso de editar' in r.get_data(as_text=True)
    with app.app_context():
        from produccion.models import LoteProduccion
        assert LoteProduccion.query.filter_by(lote='V-1').one().estado == 'abierta'


def test_alta_por_formulario_con_consumos_y_mermas(app):
    c = _login(app, 'admin')
    r = c.post('/produccion/lotes/nuevo', data={
        'producto_id': str(IDS['chorizo']), 'lote': 'W-42',
        'fecha_produccion': '2026-10-02', 'fecha_vencimiento': '',
        'peso_producido': '84,5', 'unidades_producidas': '120', 'cajas_producidas': '',
        'consumo_ingrediente_id': [str(IDS['carne']), str(IDS['sal']), str(IDS['tripa'])],
        'consumo_real': ['100', '', '60'],
        'merma_tipo': ['coccion', 'otro'],
        'merma_cantidad': ['9', '1,5'],
        'merma_motivo': ['', 'se cayó una bandeja'],
        'accion': 'guardar',
    }, follow_redirects=False)
    assert r.status_code == 302
    with app.app_context():
        from produccion.models import LoteProduccion
        from produccion import servicios
        lote = LoteProduccion.query.filter_by(lote='W-42').one()
        assert lote.peso_producido == Decimal('84.500')
        assert lote.unidades_producidas == 120 and lote.cajas_producidas is None
        assert {c.ingrediente_id for c in lote.consumos} == {IDS['carne'], IDS['tripa']}
        assert [(m.tipo, m.cantidad, m.motivo) for m in lote.mermas] == [
            ('coccion', Decimal('9.000'), None), ('otro', Decimal('1.500'), 'se cayó una bandeja')]
        assert servicios.balance(lote)['merma_sin_identificar'] == Decimal('5.000')
        assert r.headers['Location'].endswith(f'/produccion/lotes/{lote.id}')
    assert c.get(r.headers['Location']).status_code == 200


def test_alta_rechazada_conserva_lo_tecleado(app):
    c = _login(app, 'admin')
    r = c.post('/produccion/lotes/nuevo', data={
        'producto_id': str(IDS['chorizo']), 'lote': '',
        'fecha_produccion': '2026-10-02', 'peso_producido': '84',
        'consumo_ingrediente_id': [str(IDS['carne'])], 'consumo_real': ['77'],
        'merma_tipo': ['recorte'], 'merma_cantidad': ['3'], 'merma_motivo': ['borde'],
    })
    assert r.status_code == 200
    html = r.get_data(as_text=True)
    assert 'necesita un número' in html
    assert 'value="77"' in html
    assert 'value="84"' in html
    assert 'value="borde"' in html
    with app.app_context():
        from produccion.models import LoteProduccion
        assert LoteProduccion.query.count() == 0


def test_guardar_y_cerrar_en_un_paso(app):
    c = _login(app, 'admin')
    r = c.post('/produccion/lotes/nuevo', data={
        'producto_id': str(IDS['chorizo']), 'lote': 'Z-1',
        'fecha_produccion': '2026-10-02', 'peso_producido': '80',
        'consumo_ingrediente_id': [str(IDS['carne'])], 'consumo_real': ['100'],
        'accion': 'cerrar',
    }, follow_redirects=True)
    assert r.status_code == 200
    assert 'Rendimiento 80.0 %' in r.get_data(as_text=True)
    with app.app_context():
        from produccion.models import LoteProduccion
        assert LoteProduccion.query.filter_by(lote='Z-1').one().estado == 'cerrada'


def test_guardar_y_cerrar_sin_peso_queda_abierto(app):
    c = _login(app, 'admin')
    r = c.post('/produccion/lotes/nuevo', data={
        'producto_id': str(IDS['chorizo']), 'lote': 'Z-2',
        'fecha_produccion': '2026-10-02', 'peso_producido': '',
        'consumo_ingrediente_id': [str(IDS['carne'])], 'consumo_real': ['100'],
        'accion': 'cerrar',
    }, follow_redirects=True)
    assert 'Quedó guardado pero abierto' in r.get_data(as_text=True)
    with app.app_context():
        from produccion.models import LoteProduccion
        assert LoteProduccion.query.filter_by(lote='Z-2').one().estado == 'abierta'


def test_editar_cerrar_reabrir_y_anular_por_rutas(app):
    c = _login(app, 'admin')
    with app.app_context():
        lote_id = _lote().id
    assert c.get(f'/produccion/lotes/{lote_id}/editar').status_code == 200
    r = c.post(f'/produccion/lotes/{lote_id}/editar', data={
        'producto_id': str(IDS['chorizo']), 'lote': 'L-1001',
        'fecha_produccion': '2026-10-01', 'peso_producido': '88',
        'consumo_ingrediente_id': [str(IDS['carne'])], 'consumo_real': ['100'],
    }, follow_redirects=False)
    assert r.status_code == 302
    assert c.post(f'/produccion/lotes/{lote_id}/cerrar', follow_redirects=False).status_code == 302
    with app.app_context():
        from produccion.models import LoteProduccion
        lote = _db.session.get(LoteProduccion, lote_id)
        assert lote.estado == 'cerrada' and lote.peso_producido == Decimal('88.000')
    # Cerrado: editar redirige al detalle con aviso.
    r = c.get(f'/produccion/lotes/{lote_id}/editar', follow_redirects=True)
    assert 'reabrilo' in r.get_data(as_text=True)
    assert c.get(f'/produccion/lotes/{lote_id}').status_code == 200
    r = c.post(f'/produccion/lotes/{lote_id}/reabrir', data={'motivo': ''}, follow_redirects=True)
    assert 'exige un motivo' in r.get_data(as_text=True)
    c.post(f'/produccion/lotes/{lote_id}/reabrir', data={'motivo': 'faltó una merma'})
    c.post(f'/produccion/lotes/{lote_id}/anular', data={'motivo': 'prueba'})
    with app.app_context():
        from produccion.models import LoteProduccion
        assert _db.session.get(LoteProduccion, lote_id).estado == 'anulada'
    assert c.get(f'/produccion/lotes/{lote_id}').status_code == 200


def test_listado_y_reportes_responden(app):
    c = _login(app, 'admin')
    with app.app_context():
        _cerrados()
    for url in ('/produccion', '/produccion/lotes', '/produccion/lotes?estado=abierta',
                f'/produccion/lotes?producto_id={IDS["chorizo"]}&desde=2026-10-01&hasta=2026-10-31',
                '/produccion/reportes/rendimiento',
                f'/produccion/reportes/rendimiento?producto_id={IDS["chorizo"]}',
                '/produccion/reportes/mermas',
                '/produccion/reportes/mermas?desde=2026-10-01'):
        r = c.get(url)
        assert r.status_code == 200, url
    html = c.get('/produccion/reportes/rendimiento').get_data(as_text=True)
    assert '91.7 %' in html
    html = c.get('/produccion/reportes/mermas').get_data(as_text=True)
    assert 'Cocción y ahumado' in html


def test_export_excel_del_rendimiento(app):
    c = _login(app, 'admin')
    with app.app_context():
        _cerrados()
    r = c.get('/produccion/reportes/rendimiento/export?desde=2026-10-01')
    assert r.status_code == 200
    assert r.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert r.data[:2] == b'PK'
    import io
    import openpyxl
    libro = openpyxl.load_workbook(io.BytesIO(r.data))
    assert libro.sheetnames == ['Lotes', 'Por producto', 'Consumo por ingrediente']
    filas = list(libro['Lotes'].iter_rows(values_only=True))
    assert filas[0][0] == 'Código'
    assert {f[1] for f in filas[1:]} == {'A', 'B'}
    resumen = list(libro['Por producto'].iter_rows(values_only=True))
    assert resumen[1][0] == 'Chorizo' and abs(resumen[1][4] - 91.7) < 0.01


def test_lote_inexistente_da_404(app):
    c = _login(app, 'admin')
    assert c.get('/produccion/lotes/9999').status_code == 404


def test_el_menu_muestra_produccion_a_quien_puede_leer(app):
    c = _login(app, 'vend')
    html = c.get('/produccion').get_data(as_text=True)
    assert 'href="/produccion"' in html
